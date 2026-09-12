#!/usr/bin/env python3
"""Generate camera-settings.docx from camera-settings.md.

A small, purpose-built Markdown -> Word converter that understands exactly the
constructs used in camera-settings.md (headings, a table, bullet/numbered
lists, blockquotes, horizontal rules, and inline **bold** / *italic* / `code`).

Usage:
    python docs/generate_camera_settings_docx.py

Requires python-docx (pip install python-docx). If the target .docx is open in
Word (locked), the script writes camera-settings-new.docx instead.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
SRC = HERE / "camera-settings.md"
DST = HERE / "camera-settings.docx"

# Inline token pattern: `code`, **bold**, *italic* (in that precedence order).
_INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
_CODE_GREY = RGBColor(0x55, 0x55, 0x55)


def add_runs(paragraph, text: str) -> None:
    """Append inline-formatted runs (bold/italic/code) to a paragraph."""
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = _CODE_GREY
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def flush_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    # rows[0] = header, rows[1] = separator (---), rows[2:] = body
    header = rows[0]
    body = [r for r in rows[2:]]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].text = ""
        add_runs(cell.paragraphs[0], cell_text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for r in body:
        cells = table.add_row().cells
        for i, cell_text in enumerate(r):
            if i >= len(cells):
                break
            cells[i].paragraphs[0].text = ""
            add_runs(cells[i].paragraphs[0], cell_text)


def parse_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_text: str) -> Document:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    table_rows: list[list[str]] = []
    num_re = re.compile(r"^(\d+)\.\s+(.*)$")

    for raw in md_text.splitlines():
        line = raw.rstrip("\n")

        # Table accumulation.
        if line.lstrip().startswith("|"):
            table_rows.append(parse_table_row(line))
            continue
        elif table_rows:
            flush_table(doc, table_rows)
            table_rows = []

        stripped = line.strip()

        if not stripped:
            continue
        if stripped == "---":
            continue

        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, stripped[2:].strip())
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:].strip())
            continue

        m = num_re.match(stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
            add_runs(p, f"{m.group(1)}. ")
            for run in p.runs:
                run.bold = True
            add_runs(p, m.group(2))
            continue

        # Continuation line (e.g. the *Values:* detail) or plain paragraph.
        if raw.startswith(("   ", "\t")):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_runs(p, stripped)
        else:
            add_runs(doc.add_paragraph(), stripped)

    if table_rows:
        flush_table(doc, table_rows)

    return doc


def main() -> None:
    doc = convert(SRC.read_text(encoding="utf-8"))
    try:
        doc.save(DST)
        print(f"Wrote {DST}")
    except PermissionError:
        alt = HERE / "camera-settings-new.docx"
        doc.save(alt)
        print(f"{DST} was locked (open in Word?); wrote {alt} instead.")


if __name__ == "__main__":
    main()
